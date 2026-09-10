import importlib.util,tempfile,unittest,zipfile
from pathlib import Path
try:
 from docx import Document
except ImportError:Document=None
s=importlib.util.spec_from_file_location('bold_tool',Path(__file__).resolve().parents[1]/'scripts/bold-docx-phrase.py');m=importlib.util.module_from_spec(s);s.loader.exec_module(m)
@unittest.skipUnless(Document,'optional python-docx required')
class BoldTests(unittest.TestCase):
 def setUp(self):
  self.t=tempfile.TemporaryDirectory();self.addCleanup(self.t.cleanup);self.root=Path(self.t.name);self.src=self.root/'source.docx';self.out=self.root/'result.docx'
 def fixture(self,parts):
  d=Document();p=d.add_paragraph()
  for text in parts:p.add_run(text).italic=True
  d.save(self.src);self.before=self.src.read_bytes();return p.text
 def verify(self,text,phrase):
  self.assertEqual(self.before,self.src.read_bytes());d=Document(self.out);p=d.paragraphs[0];self.assertEqual(p.text,text)
  pos=text.index(phrase);chars=[(c,bool(r.bold),r.italic) for r in p.runs for c in r.text]
  for i,(_,bold,italic) in enumerate(chars):self.assertEqual(bold,pos<=i<pos+len(phrase));self.assertTrue(italic)
 def test_same_run_surrounding_text(self):
  text=self.fixture(['prefix confirm progress suffix']);m.build(self.src,self.out,'confirm progress');self.verify(text,'confirm progress')
 def test_split_runs(self):
  text=self.fixture(['prefix con','firm pro','gress suffix']);m.build(self.src,self.out,'confirm progress');self.verify(text,'confirm progress')
 def test_multiple_matches(self):
  self.fixture(['x x']);r=m.build(self.src,self.out,'x');self.assertEqual(r['occurrences_bolded'],2);self.assertEqual([bool(r.bold) for r in Document(self.out).paragraphs[0].runs],[True,False,True])
 def test_source_cannot_overwrite(self):
  self.fixture(['x'])
  with self.assertRaises(ValueError):m.build(self.src,self.src,'x')
  self.assertEqual(self.src.read_bytes(),self.before)
 def test_existing_output_preserved(self):
  self.fixture(['x']);self.out.write_bytes(b'existing')
  with self.assertRaises(ValueError):m.build(self.src,self.out,'x')
  self.assertEqual(self.out.read_bytes(),b'existing')
 def test_absent_phrase_no_output(self):
  self.fixture(['x'])
  with self.assertRaises(ValueError):m.build(self.src,self.out,'missing')
  self.assertFalse(self.out.exists())
 def test_complex_target_rejected(self):
  d=Document();p=d.add_paragraph('x');p.add_run().add_break();d.save(self.src)
  with self.assertRaises(ValueError):m.build(self.src,self.out,'x')
  self.assertFalse(self.out.exists())
 def test_fake_document_rejected(self):
  self.src.write_text('# not a Word document')
  with self.assertRaises(Exception):m.build(self.src,self.out,'x')
  self.assertFalse(self.out.exists())
if __name__=='__main__':unittest.main()
