"""Bold exact phrases across plain Word runs, preserving other run properties.
Requires python-docx. Refuses overwrite and ambiguous/complex target paragraphs.
Not a renderer or a general document rewrite engine.
"""
import argparse,copy,hashlib,json
from pathlib import Path

def build(source,output,phrase):
    from docx import Document
    from docx.oxml.ns import qn
    source=Path(source).resolve();output=Path(output).resolve()
    if source==output or output.exists():raise ValueError('new output required; source cannot be overwritten')
    if not phrase:raise ValueError('nonempty phrase required')
    before=source.read_bytes();doc=Document(source);count=0
    for para in doc.paragraphs:
        text=para.text
        if phrase not in text:continue
        runs=list(para.runs)
        if ''.join(r.text for r in runs)!=text:raise ValueError('target paragraph contains unsupported hyperlink/field text')
        if any(c.tag not in {qn('w:r'),qn('w:pPr')} for c in para._p):raise ValueError('complex target paragraph is unsupported')
        if any(c.tag not in {qn('w:rPr'),qn('w:t')} for run in runs for c in run._r):raise ValueError('complex target run is unsupported')
        spans=[];pos=0
        while True:
            start=text.find(phrase,pos)
            if start<0:break
            spans.append((start,start+len(phrase)));pos=start+len(phrase)
        offset=0
        for run in runs:
            rt=run.text;lo=offset;hi=lo+len(rt);offset=hi
            if not rt:continue
            cuts=sorted({lo,hi}|{x for span in spans for x in span if lo<x<hi})
            fragments=[]
            for a,b in zip(cuts,cuts[1:]):
                node=copy.deepcopy(run._r)
                from docx.text.run import Run
                nr=Run(node,para);nr.text=rt[a-lo:b-lo]
                if any(start<=a and b<=end for start,end in spans):nr.bold=True
                fragments.append(node)
            for node in fragments:run._r.addprevious(node)
            para._p.remove(run._r)
        if para.text!=text:raise ValueError('text preservation failed')
        count+=len(spans)
    if not count:raise ValueError('phrase absent from supported top-level paragraphs')
    if source.read_bytes()!=before:raise ValueError('source changed while preparing output')
    output.parent.mkdir(parents=True,exist_ok=True)
    # Exclusive creation: never replace a previous candidate.
    with output.open('xb') as f:doc.save(f)
    reread=Document(output)
    if [p.text for p in reread.paragraphs]!=[p.text for p in doc.paragraphs]:raise ValueError('output readback mismatch')
    return {'output':str(output),'occurrences_bolded':count,'source_sha256':hashlib.sha256(before).hexdigest(),'output_sha256':hashlib.sha256(output.read_bytes()).hexdigest(),'source_unchanged':source.read_bytes()==before,'layout_verified':False}

def main():
    ap=argparse.ArgumentParser(description=__doc__);ap.add_argument('--source',required=True);ap.add_argument('--output',required=True);ap.add_argument('--phrase',required=True);a=ap.parse_args()
    print(json.dumps(build(a.source,a.output,a.phrase),ensure_ascii=True))
if __name__=='__main__':main()
